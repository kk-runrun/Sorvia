# -*- coding: utf-8 -*-
from __future__ import annotations

import json
import os
import re
from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from super_planner.core.config import OUTPUTS_DIR, PROJECT_ROOT, WORKBENCH_ROOT
from super_planner.core.exceptions import WorkflowSystemError
from super_planner.core.naming import build_export_filename
from super_planner.schemas.project import ExportResult, PlanningProject
from super_planner.services.research_providers.base import ProgressCallback, emit_progress


EXPORT_VERSION = "planning-export-v0.1"
SHEET_STRATEGY = "1-\u6574\u4f53\u7b56\u5212\u6846\u67b6"
SHEET_LISTING = "2-Listing\u6267\u884c"
SHEET_PARAMS = "\u9644-\u53c2\u6570"
SHEET_QA = "QA\u95ee\u9898\u6e05\u5355"
REQUIRED_SHEETS = [SHEET_STRATEGY, SHEET_LISTING, SHEET_PARAMS, SHEET_QA]


def run_export(project: PlanningProject, progress: ProgressCallback | None = None) -> None:
    try:
        engine = ExportEngine(project)
        project.export_result = engine.export(progress=progress)
    except WorkflowSystemError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise WorkflowSystemError(f"\u7b56\u5212\u7a3f\u751f\u6210\u5931\u8d25\uff0c\u8bf7\u8054\u7cfbIT\u4eba\u5458\u5904\u7406\u3002Detail: {sanitize_error(exc)}") from exc


class ExportEngine:
    def __init__(self, project: PlanningProject) -> None:
        self.project = project
        self.template_path = find_template_path()
        self.workbook = self.load_template()
        self.qa_red_marks = build_red_mark_index(project.qa_result.unresolved_issues or [])
        self.listing_cell_index: list[dict[str, Any]] = []

    def red_terms_for(self, item_id: str, language: str, *, sku_id: str = "") -> list[str]:
        marks = self.qa_red_marks.get(str(item_id or ""), [])
        result = []
        for mark in marks:
            if mark.get("language") != language:
                continue
            mark_sku = str(mark.get("sku_id") or "")
            if mark_sku and sku_id and mark_sku != sku_id:
                continue
            result.append(str(mark.get("text") or ""))
        return unique_text(result)

    def export(self, progress: ProgressCallback | None = None) -> ExportResult:
        self.validate_inputs()
        emit_progress(progress, "\u6b63\u5728\u5199\u5165\u6574\u4f53\u7b56\u5212\u6846\u67b6")
        sheets = ensure_four_sheets(self.workbook)
        self.write_strategy_sheet(sheets[SHEET_STRATEGY])

        emit_progress(progress, "\u6b63\u5728\u5199\u5165Listing\u6267\u884c\u6587\u6848")
        self.write_listing_sheet(sheets[SHEET_LISTING])

        emit_progress(progress, "\u6b63\u5728\u5199\u5165\u53c2\u6570\u53cd\u67e5\u8868")
        self.write_parameter_sheet(sheets[SHEET_PARAMS])

        emit_progress(progress, "\u6b63\u5728\u5199\u5165QA\u95ee\u9898\u6e05\u5355")
        self.write_qa_sheet(sheets[SHEET_QA])

        output_dir = OUTPUTS_DIR / self.project.project_id
        output_dir.mkdir(parents=True, exist_ok=True)
        filename = build_export_filename(self.project.project_name or self.project.user_input.keywords, datetime.now())
        output_path = output_dir / filename
        self.workbook.save(output_path)

        emit_progress(progress, "\u7b56\u5212\u7a3f\u751f\u6210\u5b8c\u6210")
        return ExportResult(
            status="completed",
            summary="\u7b56\u5212\u7a3f\u751f\u6210\u5b8c\u6210\uff0c\u53ef\u4e0b\u8f7d\u6700\u7ec8 Excel\u3002",
            export_version=EXPORT_VERSION,
            generated_at=datetime.utcnow(),
            artifact_path=str(output_path),
            artifact_filename=filename,
            artifact_type="xlsx",
            template_path=str(self.template_path) if self.template_path else "",
            template_used=bool(self.template_path),
            sheet_names=REQUIRED_SHEETS,
            download_url=f"/api/projects/{self.project.project_id}/download",
            qa_unresolved_count=len(self.project.qa_result.unresolved_issues or []),
            raw_payload={
                "export_version": EXPORT_VERSION,
                "template_used": bool(self.template_path),
                "template_path": str(self.template_path) if self.template_path else "",
                "required_sheets": REQUIRED_SHEETS,
                "listing_cell_index": self.listing_cell_index[:200],
                "qa_red_mark_count": sum(len(value) for value in self.qa_red_marks.values()),
            },
        )

    def validate_inputs(self) -> None:
        if (self.project.planning_context.status or "").strip() != "completed":
            raise WorkflowSystemError("Export requires a completed Planning Context. Please contact IT.")
        if (self.project.research_result.status or "").strip() != "completed":
            raise WorkflowSystemError("Export requires a completed ResearchResult. Please contact IT.")
        if (self.project.strategy_result.status or "").strip() != "completed":
            raise WorkflowSystemError("Export requires a completed StrategyResult. Please contact IT.")
        if (self.project.content_result.status or "").strip() not in {"completed", "completed_with_issues"}:
            raise WorkflowSystemError("Export requires a completed QA-ready ContentResult. Please contact IT.")
        if (self.project.qa_result.status or "").strip() not in {"completed", "completed_with_issues"}:
            raise WorkflowSystemError("Export requires a completed QAResult. Please contact IT.")

    def load_template(self) -> Workbook:
        if self.template_path and self.template_path.exists():
            return load_workbook(self.template_path, rich_text=True)
        workbook = Workbook()
        workbook.active.title = SHEET_STRATEGY
        workbook.create_sheet(SHEET_LISTING)
        workbook.create_sheet(SHEET_PARAMS)
        workbook.create_sheet(SHEET_QA)
        return workbook

    def write_strategy_sheet(self, ws) -> None:
        prepare_strategy_sheet(ws)
        project = self.project
        strategy = project.strategy_result
        research = project.research_result

        set_cell(ws, "A1", f"Premium Project Proposal\n{project.project_name or project.user_input.keywords}-\u4ea7\u54c1\u7b56\u5212\u65b9\u6848", style="title")
        write_sku_rows(ws, 3, project.planning_context.sku_list, max_rows=max(8, len(project.planning_context.sku_list or [])))

        set_cell(ws, "B11", stringify(strategy.user_mindset), style="body")
        set_cell(ws, "C12", stringify(strategy.target_audience), style="body")
        set_cell(ws, "C13", stringify(strategy.main_uses), style="body")
        set_cell(ws, "C14", stringify(strategy.main_scenario), style="body")
        set_cell(ws, "C15", stringify(strategy.secondary_scenario), style="body")
        set_cell(ws, "C16", "User Needs Summary / \u7528\u6237\u9700\u6c42\u603b\u7ed3", style="section")
        set_cell(ws, "C17", "Focus Points", style="header")
        set_cell(ws, "D17", "Pain Points", style="header")
        set_cell(ws, "E17", "Purchase Barriers / Opportunity", style="header")
        set_cell(ws, "C18", stringify(strategy.focus_points), style="body")
        set_cell(ws, "D18", stringify(strategy.pain_points), style="body")
        set_cell(ws, "E18", stringify(strategy.purchase_barriers), style="body")
        set_cell(ws, "C23", stringify({"competitive_landscape": strategy.competitive_landscape, "market_tier": research.market_tier}), style="body")
        set_cell(ws, "B24", stringify(strategy.competitive_strategy), style="body")
        set_cell(ws, "C24", stringify(strategy.user_purchase_reason), style="body")

        for index, point in enumerate((strategy.selling_point_ranking or [])[:5], start=26):
            set_cell(ws, f"C{index}", f"TOP {point.get('rank') or index - 25}", style="header")
            set_cell(ws, f"D{index}", point.get("selling_point", ""), style="body")
            set_cell(ws, f"E{index}", point.get("reason_for_ranking", ""), style="body")

        mapping = project.content_result.content_mapping or {}
        module_lines = []
        for item in mapping.get("module_responsibilities", []) if isinstance(mapping, dict) else []:
            module_lines.append(f"{item.get('module', '')}: {item.get('role', '')}")
        set_cell(ws, "D31", "\n".join(module_lines[:2]), style="body")
        set_cell(ws, "D32", "\n".join(module_lines[2:4]), style="body")
        set_cell(ws, "D33", "\n".join(module_lines[4:]), style="body")
        set_cell(ws, "D38", stringify(project.content_result.content_mapping.get("selling_point_distribution", [])), style="body")
        set_cell(ws, "D39", stringify(strategy.communication_strategy), style="body")
        set_cell(ws, "D40", stringify(strategy.competitive_opportunity), style="body")
        set_cell(ws, "D41", stringify(strategy.needs_verification), style="body")

        row = 45
        row = write_section_table(ws, row, "Product / Market Judgment", [{"Field": "Product Category", "Content": stringify(strategy.product_category), "Evidence / Confidence": stringify(strategy.confidence)}])
        row = write_section_table(ws, row, "Communication Strategy", [{"Field": "Main Line", "Content": stringify(strategy.communication_strategy), "Evidence / Confidence": stringify(strategy.communication_strategy.get("evidence_ids", []))}])
        rows = []
        for point in (strategy.selling_point_ranking or [])[:5]:
            rows.append(
                {
                    "Rank": f"TOP {point.get('rank', '')}",
                    "Selling Point": point.get("selling_point", ""),
                    "User Need": point.get("corresponding_user_need", ""),
                    "Product Fact": point.get("corresponding_product_fact", ""),
                    "Market Basis": point.get("corresponding_market_competitor_basis", ""),
                    "Ranking Reason": point.get("reason_for_ranking", ""),
                    "Evidence / Confidence": f"{', '.join(point.get('evidence_ids', []) or [])}\nconfidence={point.get('confidence', '')}",
                }
            )
        write_table(ws, row, ["Rank", "Selling Point", "User Need", "Product Fact", "Market Basis", "Ranking Reason", "Evidence / Confidence"], rows)
        autosize_columns(ws, max_width=46)

    def write_listing_sheet(self, ws) -> None:
        clear_sheet(ws)
        prepare_listing_sheet(ws)
        project = self.project
        items = project.content_result.content_items or []
        sku_list = project.planning_context.sku_list or []

        row = 1
        row = write_title(ws, row, "Overview of All SKUs / \u5168SKU\u6982\u89c8")
        row = write_table(
            ws,
            row,
            ["No.", "SKU", "Name"],
            [{"No.": sku.sequence, "SKU": sku.sku, "Name": sku.name} for sku in sku_list],
        )
        row += 1

        title_rows = []
        for sku in sku_list:
            title_rows.append(
                {
                    "No.": sku.sequence,
                    "SKU": sku.sku,
                    "Item Name": text_for_item(find_content_item(items, "Title", "Item Name"), "english"),
                    "Item Highlight": text_for_item(find_content_item(items, "Title Reference", "Item Highlight"), "english"),
                    "Item Name CN": text_for_item(find_content_item(items, "Title", "Item Name"), "chinese"),
                    "Item Highlight CN": text_for_item(find_content_item(items, "Title Reference", "Item Highlight"), "chinese"),
                }
            )
        row = write_title(ws, row, "Title Reference / \u6807\u9898\u53c2\u8003")
        title_header_row = row
        row = write_table(ws, row, ["No.", "SKU", "Item Name", "Item Highlight", "Item Name CN", "Item Highlight CN"], title_rows)
        self.apply_title_red_marks(ws, title_header_row + 1, len(title_rows))
        row += 1

        row = write_title(ws, row, "Supporting Images / \u56fe\u7247\u6587\u6848")
        supporting_headers = ["Section", "Copy Content EN", "Copy Content CN", "SKU Scope", "Parameter Ref", "Evidence Ref", "Confidence"]
        supporting_rows = []
        for item in sorted([item for item in items if item.get("module") == "Supporting Image Copy"], key=lambda value: int(value.get("sequence") or 999)):
            supporting_rows.append(content_row(item, item.get("submodule", "")))
            for variant in item.get("sku_copy_variants") or []:
                if isinstance(variant, dict):
                    supporting_rows.append(
                        {
                            "Section": f"{item.get('submodule', '')} / {variant.get('sku_id', '')}",
                            "Copy Content EN": variant.get("english_copy", ""),
                            "Copy Content CN": variant.get("chinese_copy", ""),
                            "SKU Scope": "SKU_SPECIFIC",
                            "Parameter Ref": stringify(variant.get("parameter_reference", [])),
                            "Evidence Ref": stringify(item.get("evidence_reference", [])),
                            "Confidence": item.get("confidence", ""),
                            "_item_id": item.get("id", ""),
                            "_sku_id": variant.get("sku_id", ""),
                        }
                    )
        start_row = row + 1
        row = write_table(ws, row, supporting_headers, supporting_rows, export_engine=self)
        row += 1

        row = write_title(ws, row, "Premium A+ / \u9ad8\u7ea7A+\u6587\u6848")
        aplus_headers = ["Section", "Copy Content EN", "Copy Content CN", "SKU Scope", "Parameter Ref", "Evidence Ref", "Confidence"]
        aplus_rows = [content_row(item, f"Module {item.get('sequence')} / {item.get('submodule', '')}") for item in sorted([item for item in items if item.get("module") == "Premium A+"], key=lambda value: int(value.get("sequence") or 999))]
        row = write_table(ws, row, aplus_headers, aplus_rows, export_engine=self)
        row += 1

        row = write_title(ws, row, "Bullet Point / \u4e94\u70b9\u63cf\u8ff0")
        bullet_rows = [content_row(item, str(item.get("sequence", ""))) for item in sorted([item for item in items if item.get("module") == "Bullet Point"], key=lambda value: int(value.get("sequence") or 999))]
        row = write_table(ws, row, aplus_headers, bullet_rows, export_engine=self)
        row += 1

        row = write_title(ws, row, "Product Description / \u4ea7\u54c1\u63cf\u8ff0")
        desc_rows = [content_row(item, item.get("submodule", "")) for item in sorted([item for item in items if item.get("module") == "Product Description"], key=lambda value: int(value.get("sequence") or 999))]
        row = write_table(ws, row, aplus_headers, desc_rows, export_engine=self)
        autosize_columns(ws, max_width=58)
        ws.freeze_panes = "A2"

    def apply_title_red_marks(self, ws, start_row: int, row_count: int) -> None:
        item_name_id = "content-title-item-name-1"
        item_highlight_id = "content-title-reference-item-highlight-1"
        for row in range(start_row, start_row + row_count):
            set_cell_rich(ws.cell(row=row, column=3), stringify(ws.cell(row=row, column=3).value), self.red_terms_for(item_name_id, "english"))
            set_cell_rich(ws.cell(row=row, column=4), stringify(ws.cell(row=row, column=4).value), self.red_terms_for(item_highlight_id, "english"))
            set_cell_rich(ws.cell(row=row, column=5), stringify(ws.cell(row=row, column=5).value), self.red_terms_for(item_name_id, "chinese"))
            set_cell_rich(ws.cell(row=row, column=6), stringify(ws.cell(row=row, column=6).value), self.red_terms_for(item_highlight_id, "chinese"))
            self.listing_cell_index.extend(
                [
                    {"item_id": item_name_id, "field": "english_copy", "cell": f"C{row}"},
                    {"item_id": item_highlight_id, "field": "english_copy", "cell": f"D{row}"},
                    {"item_id": item_name_id, "field": "chinese_copy", "cell": f"E{row}"},
                    {"item_id": item_highlight_id, "field": "chinese_copy", "cell": f"F{row}"},
                ]
            )

    def write_parameter_sheet(self, ws) -> None:
        clear_sheet(ws)
        prepare_plain_sheet(ws)
        project = self.project
        row = write_title(ws, 1, "\u9644-\u53c2\u6570 / Task 2 \u89c4\u683c\u4e8b\u5b9e\u53cd\u67e5")
        row = write_table(
            ws,
            row,
            ["No.", "SKU ID", "SKU", "Name", "Source File"],
            [
                {
                    "No.": sku.sequence,
                    "SKU ID": sku.sku_id,
                    "SKU": sku.sku,
                    "Name": sku.name,
                    "Source File": sku.source_file_name,
                }
                for sku in project.planning_context.sku_list or []
            ],
        )
        row += 1

        row = write_title(ws, row, "\u6807\u51c6\u4e2d\u6587\u53c2\u6570\u8868")
        sku_ids = [sku.sku_id for sku in project.planning_context.sku_list or []]
        sku_labels = [sku.sku or sku.sku_id for sku in project.planning_context.sku_list or []]
        headers = ["Parameter", "Common Values"] + sku_labels
        rows = []
        for item in project.planning_context.standard_chinese_parameter_table or []:
            values_by_sku = getattr(item, "values_by_sku", {}) or {}
            rows.append(
                {
                    "Parameter": getattr(item, "parameter", ""),
                    "Common Values": " / ".join([str(value) for value in getattr(item, "values", []) if str(value).strip()]),
                    **{label: values_by_sku.get(sku_id, "") for sku_id, label in zip(sku_ids, sku_labels)},
                }
            )
        row = write_table(ws, row, headers, rows)
        row += 1

        row = write_title(ws, row, "\u4ea7\u54c1\u516c\u5171\u53c2\u6570")
        row = write_table(
            ws,
            row,
            ["Parameter", "Value"],
            [{"Parameter": key, "Value": value} for key, value in (project.planning_context.product_common_parameters or {}).items()],
        )
        row += 1

        row = write_title(ws, row, "SKU Group \u516c\u5171\u53c2\u6570")
        group_rows = []
        for group in project.planning_context.sku_group_common_parameters or []:
            group_rows.append(
                {
                    "Group ID": group.group_id,
                    "SKU IDs": ", ".join(group.sku_ids or []),
                    "SKU Indexes": ", ".join([str(index) for index in group.sku_indexes or []]),
                    "Parameters": stringify(group.parameters),
                }
            )
        row = write_table(ws, row, ["Group ID", "SKU IDs", "SKU Indexes", "Parameters"], group_rows)
        row += 1

        row = write_title(ws, row, "\u5355SKU\u5dee\u5f02\u53c2\u6570")
        diff_rows = []
        for sku_id, params in (project.planning_context.sku_differential_parameters or {}).items():
            for key, value in (params or {}).items():
                diff_rows.append({"SKU ID": sku_id, "Parameter": key, "Value": value})
        write_table(ws, row, ["SKU ID", "Parameter", "Value"], diff_rows)
        autosize_columns(ws, max_width=52)
        ws.freeze_panes = "A3"

    def write_qa_sheet(self, ws) -> None:
        clear_sheet(ws)
        prepare_plain_sheet(ws)
        row = write_title(ws, 1, "QA\u95ee\u9898\u6e05\u5355")
        unresolved = [issue for issue in self.project.qa_result.unresolved_issues or [] if issue.get("status") == "unresolved"]
        if not unresolved:
            set_cell(ws, f"A{row}", "\u672c\u6b21QA\u672a\u53d1\u73b0\u9700\u8981\u4eba\u5de5\u5904\u7406\u7684\u95ee\u9898", style="body")
            autosize_columns(ws, max_width=42)
            return
        headers = [
            "SKU Scope",
            "SKU ID",
            "Module",
            "Issue Type",
            "Risk Level",
            "Problem Text / Position",
            "Final English Copy",
            "Final Chinese Copy",
            "Issue Description",
            "Rewrite Attempts",
            "Suggested Action",
            "Red Mark Spans",
        ]
        rows = []
        for issue in unresolved:
            rows.append(
                {
                    "SKU Scope": issue.get("sku_scope", ""),
                    "SKU ID": issue.get("sku_id", ""),
                    "Module": f"{issue.get('module', '')} / {issue.get('submodule', '')}",
                    "Issue Type": issue.get("issue_type", ""),
                    "Risk Level": issue.get("risk_level", ""),
                    "Problem Text / Position": stringify(issue.get("problematic_span") or issue.get("problematic_text", "")),
                    "Final English Copy": issue.get("english_copy", ""),
                    "Final Chinese Copy": issue.get("chinese_copy", ""),
                    "Issue Description": issue.get("issue_description", ""),
                    "Rewrite Attempts": str(len(issue.get("rewrite_attempts") or [])),
                    "Suggested Action": issue.get("suggested_action", ""),
                    "Red Mark Spans": stringify(issue.get("red_mark_spans", [])),
                }
            )
        write_table(ws, row, headers, rows, risk_col="Risk Level")
        autosize_columns(ws, max_width=54)
        ws.freeze_panes = "A3"


def find_template_path() -> Path | None:
    explicit = os.getenv("SUPER_PLANNER_EXPORT_TEMPLATE", "").strip()
    candidates = [Path(explicit)] if explicit else []
    candidates.extend(
        [
            PROJECT_ROOT / "\u7b56\u5212\u7a3f\u6a21\u677f.xlsx",
            PROJECT_ROOT / "templates" / "\u7b56\u5212\u7a3f\u6a21\u677f.xlsx",
            PROJECT_ROOT.parent / "\u7b56\u5212\u7a3f\u6a21\u677f.xlsx",
            WORKBENCH_ROOT / "\u7b56\u5212\u7a3f\u6a21\u677f.xlsx",
        ]
    )
    for path in candidates:
        if path and path.exists():
            return path
    return None


def ensure_four_sheets(workbook: Workbook) -> dict[str, Any]:
    sheets = {}
    for index, title in enumerate(REQUIRED_SHEETS):
        if title in workbook.sheetnames:
            ws = workbook[title]
        elif index < len(workbook.worksheets):
            ws = workbook.worksheets[index]
            ws.title = title
        else:
            ws = workbook.create_sheet(title)
        sheets[title] = ws
    for ws in list(workbook.worksheets):
        if ws.title not in REQUIRED_SHEETS:
            workbook.remove(ws)
    workbook._sheets = [sheets[title] for title in REQUIRED_SHEETS]  # keep required order
    return sheets


def prepare_strategy_sheet(ws) -> None:
    ws.sheet_view.showGridLines = False
    for col, width in {"A": 20, "B": 28, "C": 30, "D": 40, "E": 52, "F": 28, "G": 28}.items():
        ws.column_dimensions[col].width = width
    for row in range(1, max(ws.max_row, 80) + 1):
        ws.row_dimensions[row].height = 42


def prepare_listing_sheet(ws) -> None:
    ws.sheet_view.showGridLines = False
    widths = {"A": 18, "B": 42, "C": 42, "D": 20, "E": 34, "F": 34, "G": 16}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def prepare_plain_sheet(ws) -> None:
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 32
    ws.column_dimensions["C"].width = 38
    ws.column_dimensions["D"].width = 38


def clear_sheet(ws) -> None:
    for merged in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merged))
    for row in ws.iter_rows():
        for cell in row:
            cell.value = None


def write_title(ws, row: int, title: str) -> int:
    cell = ws.cell(row=row, column=1)
    cell.value = title
    apply_style(cell, "section")
    for column in range(1, 8):
        apply_border(ws.cell(row=row, column=column))
    ws.row_dimensions[row].height = 28
    return row + 1


def write_section_table(ws, row: int, title: str, rows: list[dict[str, Any]]) -> int:
    row = write_title(ws, row, title)
    headers = list(rows[0].keys()) if rows else ["Field", "Content"]
    return write_table(ws, row, headers, rows) + 1


def write_table(
    ws,
    row: int,
    headers: list[str],
    rows: list[dict[str, Any]],
    *,
    export_engine: ExportEngine | None = None,
    risk_col: str | None = None,
) -> int:
    for column, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=column)
        cell.value = header
        apply_style(cell, "header")
    row += 1
    if not rows:
        cell = ws.cell(row=row, column=1)
        cell.value = "\u6682\u65e0\u6570\u636e"
        apply_style(cell, "body")
        return row + 1
    for data in rows:
        item_id = str(data.get("_item_id") or "")
        sku_id = str(data.get("_sku_id") or "")
        for column, header in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=column)
            value = data.get(header, "")
            if export_engine and header == "Copy Content EN":
                red_terms = export_engine.red_terms_for(item_id, "english", sku_id=sku_id)
                set_cell_rich(cell, stringify(value), red_terms)
                export_engine.listing_cell_index.append({"item_id": item_id, "sku_id": sku_id, "field": "english_copy", "cell": cell.coordinate})
            elif export_engine and header == "Copy Content CN":
                red_terms = export_engine.red_terms_for(item_id, "chinese", sku_id=sku_id)
                set_cell_rich(cell, stringify(value), red_terms)
                export_engine.listing_cell_index.append({"item_id": item_id, "sku_id": sku_id, "field": "chinese_copy", "cell": cell.coordinate})
            else:
                cell.value = stringify(value)
            apply_style(cell, "body")
            if risk_col and header == risk_col:
                apply_risk_style(cell, stringify(value))
        ws.row_dimensions[row].height = max(24, min(120, 18 + max(len(str(data.get(header, ""))) for header in headers) // 6))
        row += 1
    return row


def register_listing_item_cells(export_engine: ExportEngine, item_id: str, start_row: int, english_col: int, chinese_col: int) -> None:
    for row in range(start_row, start_row + max(1, len(export_engine.project.planning_context.sku_list or []))):
        export_engine.listing_cell_index.append({"item_id": item_id, "field": "english_copy", "cell": f"{get_column_letter(english_col)}{row}"})
        export_engine.listing_cell_index.append({"item_id": item_id, "field": "chinese_copy", "cell": f"{get_column_letter(chinese_col)}{row}"})


def write_sku_rows(ws, start_row: int, sku_list: list[Any], max_rows: int) -> None:
    for offset in range(max_rows):
        row = start_row + offset
        sku = sku_list[offset] if offset < len(sku_list) else None
        if not sku:
            continue
        set_cell(ws, f"C{row}", sku.sequence, style="body")
        set_cell(ws, f"D{row}", sku.sku, style="body")
        set_cell(ws, f"E{row}", sku.name, style="body")


def content_row(item: dict[str, Any], section: str) -> dict[str, Any]:
    return {
        "Section": section,
        "Copy Content EN": text_for_item(item, "english"),
        "Copy Content CN": text_for_item(item, "chinese"),
        "SKU Scope": item.get("sku_scope", ""),
        "Parameter Ref": stringify(item.get("parameter_reference", [])),
        "Evidence Ref": stringify(item.get("evidence_reference", [])),
        "Confidence": item.get("confidence", ""),
        "_item_id": item.get("id", ""),
    }


def text_for_item(item: dict[str, Any] | None, language: str) -> str:
    if not item:
        return ""
    key = f"{language}_copy"
    if isinstance(item.get("copy_parts"), dict):
        order = item.get("copy_part_order") or list(item["copy_parts"].keys())
        values = []
        for name in order:
            part = item["copy_parts"].get(name)
            if isinstance(part, dict) and str(part.get(key) or "").strip():
                values.append(str(part.get(key)).strip())
        return "\n".join(values).strip()
    return str(item.get(key) or "").strip()


def find_content_item(items: list[dict[str, Any]], module: str, submodule: str) -> dict[str, Any] | None:
    for item in items:
        if item.get("module") == module and item.get("submodule") == submodule:
            return item
    return None


def set_cell(ws, coordinate: str, value: Any, *, style: str = "body") -> None:
    cell = writable_cell(ws, coordinate)
    cell.value = stringify(value)
    apply_style(cell, style)


def writable_cell(ws, coordinate: str):
    for merged in ws.merged_cells.ranges:
        if coordinate in merged:
            return ws.cell(row=merged.min_row, column=merged.min_col)
    return ws[coordinate]


def set_cell_rich(cell, text: str, red_terms: list[str]) -> None:
    text = str(text or "")
    terms = [term for term in unique_text(red_terms) if term and term.lower() in text.lower()]
    if not terms:
        cell.value = text
        return
    try:
        cell.value = build_rich_text(text, terms)
    except Exception:
        cell.value = text
        cell.font = copy(cell.font)
        cell.font = Font(name=cell.font.name or "Calibri", size=cell.font.sz or 11, color="FFFF0000")


def build_rich_text(text: str, red_terms: list[str]) -> CellRichText:
    spans = []
    occupied = [False] * len(text)
    lower = text.lower()
    for term in sorted(red_terms, key=len, reverse=True):
        start = lower.find(term.lower())
        while start >= 0:
            end = start + len(term)
            if not any(occupied[start:end]):
                spans.append((start, end))
                for index in range(start, end):
                    occupied[index] = True
            start = lower.find(term.lower(), end)
    spans.sort()
    rich = CellRichText()
    position = 0
    red_font = InlineFont(color="FFFF0000")
    for start, end in spans:
        if start > position:
            rich.append(text[position:start])
        rich.append(TextBlock(red_font, text[start:end]))
        position = end
    if position < len(text):
        rich.append(text[position:])
    return rich


def apply_style(cell, style: str) -> None:
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    apply_border(cell)
    if style == "title":
        cell.font = Font(bold=True, size=14, color="FFFFFFFF")
        cell.fill = PatternFill("solid", fgColor="FF1F2937")
        return
    if style == "section":
        cell.font = Font(bold=True, size=12, color="FFFFFFFF")
        cell.fill = PatternFill("solid", fgColor="FF374151")
        return
    if style == "header":
        cell.font = Font(bold=True, color="FF111827")
        cell.fill = PatternFill("solid", fgColor="FFE5E7EB")
        return
    cell.font = Font(color="FF111827")


def apply_border(cell) -> None:
    thin = Side(style="thin", color="FFD1D5DB")
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)


def apply_risk_style(cell, risk: str) -> None:
    normalized = risk.lower()
    if normalized == "high":
        cell.fill = PatternFill("solid", fgColor="FFFFD6D6")
        cell.font = Font(bold=True, color="FFB91C1C")
    elif normalized == "medium":
        cell.fill = PatternFill("solid", fgColor="FFFFF1C2")
        cell.font = Font(bold=True, color="FF92400E")
    elif normalized == "low":
        cell.fill = PatternFill("solid", fgColor="FFE0F2FE")
        cell.font = Font(bold=True, color="FF075985")


def autosize_columns(ws, *, max_width: int = 60) -> None:
    for column_cells in ws.columns:
        letter = get_column_letter(column_cells[0].column)
        current = ws.column_dimensions[letter].width or 10
        max_len = current
        for cell in column_cells[:120]:
            value = cell.value
            if value is None:
                continue
            text = str(value)
            lines = text.splitlines() or [text]
            max_len = max(max_len, min(max_width, max(len(line) for line in lines) + 2))
        ws.column_dimensions[letter].width = min(max_width, max_len)


def build_red_mark_index(unresolved_issues: list[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    index: dict[str, list[dict[str, Any]]] = {}
    for issue in unresolved_issues:
        if issue.get("status") != "unresolved":
            continue
        item_id = str(issue.get("content_item_id") or "")
        if not item_id:
            continue
        for span in issue.get("red_mark_spans") or []:
            if not isinstance(span, dict):
                continue
            text = str(span.get("text") or issue.get("problematic_text") or "").strip()
            if not text:
                continue
            field = str(span.get("field") or "")
            language = "chinese" if "chinese" in field.lower() else "english"
            index.setdefault(item_id, []).append(
                {
                    "text": text,
                    "language": language,
                    "field": field,
                    "sku_id": issue.get("sku_id", ""),
                    "issue_type": issue.get("issue_type", ""),
                    "risk_level": issue.get("risk_level", ""),
                }
            )
    return index


def register_red_marks(export_engine: ExportEngine, item_id: str, english_cell, chinese_cell, sku_id: str = "") -> None:
    export_engine.listing_cell_index.append({"item_id": item_id, "sku_id": sku_id, "field": "english_copy", "cell": english_cell.coordinate})
    export_engine.listing_cell_index.append({"item_id": item_id, "sku_id": sku_id, "field": "chinese_copy", "cell": chinese_cell.coordinate})


def stringify(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, (int, float, bool)):
        return str(value)
    if hasattr(value, "model_dump"):
        value = value.model_dump(mode="json")
    if isinstance(value, dict):
        compact = []
        for key, item in value.items():
            if item in ("", None) or item == [] or item == {}:
                continue
            compact.append(f"{key}: {stringify(item)}")
        return "\n".join(compact)
    if isinstance(value, list):
        lines = []
        for index, item in enumerate(value, start=1):
            text = stringify(item)
            if text:
                lines.append(f"{index}. {text}")
        return "\n".join(lines)
    return str(value)


def unique_text(values: list[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        key = value.lower()
        if not key or key in seen:
            continue
        seen.add(key)
        result.append(value)
    return result


def sanitize_error(exc: Exception) -> str:
    text = str(exc)
    text = re.sub(r"sk-[A-Za-z0-9_\-]+", "sk-***", text)
    text = re.sub(r"Bearer\s+[A-Za-z0-9_\-.]+", "Bearer ***", text, flags=re.IGNORECASE)
    return re.sub(r"\s+", " ", text).strip()[:500]
